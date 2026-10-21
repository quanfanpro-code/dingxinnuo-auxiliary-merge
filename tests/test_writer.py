"""结果工作簿写入器的最小行为检查。"""

import io
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree
import importlib.util

from openpyxl import load_workbook


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "写入结果工作簿.py"
SPEC = importlib.util.spec_from_file_location("result_writer", SCRIPT)
writer = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(writer)


class WriterTest(unittest.TestCase):
    def test_多表类型和版式一次写成(self):
        output = io.BytesIO()
        sheets = [
            {
                "名称": "序时账合并",
                "列": [
                    {"名称": "科目编号", "类型": "文本", "宽度": 16},
                    {"名称": "记账日期", "类型": "日期", "宽度": 15},
                    {"名称": "借方发生额", "类型": "金额", "宽度": 18},
                    {"名称": "辅助核算信息", "类型": "文本", "宽度": 28},
                ],
                "行": iter([
                    ["0012", "2026-10-21", "1234.50", "客户：甲公司"],
                    ["合计", "", "1234.50", "=HYPERLINK(\"x\")"],
                ]),
            },
            {"名称": "核对结果", "列": [{"名称": "事项", "类型": "文本"}], "行": iter([])},
        ]
        self.assertEqual(writer.write_workbook(sheets, output), {"序时账合并": 2, "核对结果": 0})
        output.seek(0)
        with ZipFile(output) as archive:
            for member in ("xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml"):
                ElementTree.fromstring(archive.read(member))
        output.seek(0)
        book = load_workbook(output)
        sheet = book["序时账合并"]
        self.assertEqual(sheet["B2"].value, "0012")
        self.assertEqual(sheet["C2"].value.date(), date(2026, 10, 21))
        self.assertEqual(sheet["D2"].value, 1234.5)
        self.assertEqual(sheet["E3"].value, '=HYPERLINK("x")')
        self.assertEqual(sheet["E3"].data_type, "s")
        self.assertEqual(sheet["B1"].fill.fgColor.rgb, "001F4E79")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["D2"].number_format, "#,##0.00")
        self.assertEqual(book["核对结果"].max_row, 1)
        book.close()

    def test_超出_excel_数值精度时拒绝(self):
        output = io.BytesIO()
        sheets = [{"名称": "余额表", "列": [{"名称": "期末余额", "类型": "金额"}],
                   "行": [[Decimal("1234567890123456")]]}]
        with self.assertRaisesRegex(ValueError, "15 位"):
            writer.write_workbook(sheets, output)


if __name__ == "__main__":
    unittest.main()
